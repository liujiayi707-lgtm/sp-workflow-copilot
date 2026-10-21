
"""Local demo portal backed by anonymized SocialPeta exports."""

import csv
import io
from datetime import date, datetime
from html import escape
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import parse_qs, quote, unquote, urlparse

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[3]
INPUT_DIR = ROOT / "data" / "flow_a_demo"
HOST = "127.0.0.1"
PORT = 8766

SALES_MAP = {
    "公司A": "Sales A",
    "公司B": "Sales B",
    "公司C": "Sales C",
}

SLUG_MAP = {
    "公司A": "company-a",
    "公司B": "company-b",
    "公司C": "company-c",
}


def parse_date(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, str):
        try:
            return datetime.strptime(value.strip(), "%Y-%m-%d").strftime(
                "%Y-%m-%d"
            )
        except ValueError:
            return None
    return None


def load_companies():
    companies = {}

    for path in sorted(INPUT_DIR.glob("*.xlsx")):
        wb = load_workbook(path, read_only=True, data_only=True)

        try:
            ws = wb["访问数据"]
            metadata = {}

            for row in ws.iter_rows(min_row=1, max_row=6, values_only=True):
                if len(row) >= 2 and row[0] is not None:
                    metadata[str(row[0]).strip()] = row[1]

            company_name = str(metadata.get("公司名", "")).strip()

            if company_name not in SALES_MAP:
                raise ValueError(
                    f"文件 {path.name} 的公司名无法匹配 Sales："
                    f"{company_name!r}"
                )

            header_row = None
            headers = []

            for row_num, row in enumerate(
                ws.iter_rows(values_only=True), start=1
            ):
                values = list(row)
                if values and str(values[0]).strip() == "用户id":
                    header_row = row_num
                    headers = values
                    break

            if header_row is None:
                raise ValueError(f"文件缺少用户表头：{path.name}")

            positions = {
                str(value).strip(): index
                for index, value in enumerate(headers)
                if value is not None
            }

            user_col = positions["用户id"]
            label_col = positions["高低频率"]

            date_columns = []
            for index, value in enumerate(headers):
                day = parse_date(value)
                if day:
                    date_columns.append((index, day))

            accounts = []

            for row in ws.iter_rows(
                min_row=header_row + 1,
                values_only=True,
            ):
                values = list(row)

                if len(values) <= user_col or values[user_col] is None:
                    continue

                user_id = str(values[user_col]).strip()
                if not user_id:
                    continue

                label_value = (
                    values[label_col]
                    if label_col < len(values)
                    else None
                )

                usage_by_date = {}

                for col_index, day in date_columns:
                    raw = (
                        values[col_index]
                        if col_index < len(values)
                        else None
                    )

                    if raw is None or (
                        isinstance(raw, str) and not raw.strip()
                    ):
                        usage_by_date[day] = None
                    elif (
                        isinstance(raw, bool)
                        or not isinstance(raw, (int, float))
                        or raw < 0
                        or (
                            isinstance(raw, float)
                            and not raw.is_integer()
                        )
                    ):
                        usage_by_date[day] = None
                    else:
                        usage_by_date[day] = int(raw)

                valid = [
                    value
                    for value in usage_by_date.values()
                    if value is not None
                ]

                accounts.append({
                    "user_id": user_id,
                    "frequency_label": (
                        str(label_value).strip()
                        if label_value is not None
                        else ""
                    ),
                    "usage_by_date": usage_by_date,
                    "total_usage": sum(valid) if valid else None,
                    "valid_days": len(valid),
                })

            companies[company_name] = {
                "name": company_name,
                "sales": SALES_MAP[company_name],
                "slug": SLUG_MAP[company_name],
                "dates": [day for _, day in date_columns],
                "accounts": accounts,
                "source_file": path.name,
            }

        finally:
            wb.close()

    if set(companies) != set(SALES_MAP):
        raise ValueError(
            "必须且只能找到公司 A、公司 B、公司 C 三份测试文件。"
        )

    return companies


COMPANIES = load_companies()


def html_page(title, body):
    return f"""<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{escape(title)}</title>
<style>
body {{ font-family: -apple-system, sans-serif; max-width: 1000px;
       margin: 36px auto; padding: 0 20px; color: #222; }}
table {{ border-collapse: collapse; width: 100%; margin-top: 18px; }}
th, td {{ border: 1px solid #ddd; padding: 9px; text-align: left; }}
th {{ background: #f3f4f6; }}
.notice {{ background: #f3f4f6; padding: 12px; border-radius: 8px; }}
button, select {{ padding: 7px; }}
</style>
</head>
<body>
<h1>{escape(title)}</h1>
{body}
</body>
</html>"""


class Handler(BaseHTTPRequestHandler):
    def send_html(self, body, status=200):
        data = body.encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self):
        parsed = urlparse(self.path)
        path = parsed.path
        query = parse_qs(parsed.query)

        if path == "/":
            selected = query.get("business", ["All"])[0]
            options = '<option value="All">全部商务</option>'

            for sales in SALES_MAP.values():
                mark = " selected" if selected == sales else ""
                options += (
                    f'<option value="{escape(sales)}"{mark}>'
                    f'{escape(sales)}</option>'
                )

            rows = ""

            for company in COMPANIES.values():
                if selected not in ("All", company["sales"]):
                    continue

                high = sum(
                    account["frequency_label"] == "高频"
                    for account in company["accounts"]
                )
                low = sum(
                    account["frequency_label"] == "低频"
                    for account in company["accounts"]
                )

                rows += f"""
                <tr>
                  <td>{escape(company['sales'])}</td>
                  <td><a href="/customer/{company['slug']}">
                    {escape(company['name'])}</a></td>
                  <td>{len(company['accounts'])}</td>
                  <td>{high}</td>
                  <td>{low}</td>
                  <td>{len(company['dates'])}</td>
                </tr>"""

            body = f"""
            <p class="notice">
              本地测试环境：使用脱敏导出数据，不连接真实 SocialPeta。
            </p>
            <form action="/" method="get">
              <label for="business">负责商务：</label>
              <select id="business" name="business">{options}</select>
              <button type="submit">筛选客户</button>
            </form>
            <table>
              <thead><tr>
                <th>负责商务</th><th>公司</th><th>账号数</th>
                <th>高频账号</th><th>低频账号</th><th>日期列数</th>
              </tr></thead>
              <tbody>{rows}</tbody>
            </table>"""

            return self.send_html(html_page("Flow A 真实数据测试后台", body))

        if path.startswith("/customer/"):
            slug = path.rsplit("/", 1)[-1]
            company = next(
                (item for item in COMPANIES.values()
                 if item["slug"] == slug),
                None,
            )

            if not company:
                return self.send_html(
                    html_page("未找到客户", "<p>公司不存在。</p>"),
                    404,
                )

            rows = ""

            for account in company["accounts"]:
                total = account["total_usage"]
                total_text = str(total) if total is not None else "数据缺失"

                rows += f"""
                <tr>
                  <td>{escape(account['user_id'])}</td>
                  <td>{escape(account['frequency_label'])}</td>
                  <td>{total_text}</td>
                  <td>{account['valid_days']}</td>
                </tr>"""

            body = f"""
            <p><strong>负责商务：</strong>{escape(company['sales'])}</p>
            <p><strong>公司：</strong>{escape(company['name'])}</p>
            <p><strong>日期列数：</strong>{len(company['dates'])}</p>
            <table>
              <thead><tr>
                <th>用户 ID</th><th>原平台高低频标签</th>
                <th>累计使用量</th><th>有效日期数</th>
              </tr></thead>
              <tbody>{rows}</tbody>
            </table>
            <p><a href="/download/{company['slug']}.csv"
                  download>下载使用数据 CSV</a></p>
            <p><a href="/">返回客户列表</a></p>"""

            return self.send_html(
                html_page(f"{company['name']} 使用数据", body)
            )

        if path.startswith("/download/") and path.endswith(".csv"):
            slug = path.rsplit("/", 1)[-1][:-4]
            company = next(
                (item for item in COMPANIES.values()
                 if item["slug"] == slug),
                None,
            )

            if not company:
                return self.send_html(
                    html_page("下载失败", "<p>公司不存在。</p>"),
                    404,
                )

            output = io.StringIO(newline="")
            writer = csv.writer(output)
            writer.writerow([
                "company_name",
                "sales",
                "user_id",
                "frequency_label",
                *company["dates"],
            ])

            for account in company["accounts"]:
                writer.writerow([
                    company["name"],
                    company["sales"],
                    account["user_id"],
                    account["frequency_label"],
                    *[
                        account["usage_by_date"][day]
                        if account["usage_by_date"][day] is not None
                        else ""
                        for day in company["dates"]
                    ],
                ])

            data = output.getvalue().encode("utf-8-sig")
            filename = f"{slug}.csv"

            self.send_response(200)
            self.send_header("Content-Type", "text/csv; charset=utf-8")
            self.send_header(
                "Content-Disposition",
                f'attachment; filename="{filename}"',
            )
            self.send_header("Content-Length", str(len(data)))
            self.end_headers()
            self.wfile.write(data)
            return

        self.send_html(html_page("未找到页面", "<p>404 Not Found</p>"), 404)

    def log_message(self, format_string, *args):
        print("[真实数据模拟后台]", format_string % args)


if __name__ == "__main__":
    server = ThreadingHTTPServer((HOST, PORT), Handler)
    print(f"真实数据模拟后台：http://{HOST}:{PORT}")
    for company in COMPANIES.values():
        print(
            f"{company['sales']} -> {company['name']}: "
            f"{len(company['accounts'])} 个账号"
        )
    print("按 Control + C 停止服务。")

    try:
        server.serve_forever()
    except KeyboardInterrupt:
        print("\n模拟后台已停止。")
    finally:
        server.server_close()
