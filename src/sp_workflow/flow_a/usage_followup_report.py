
"""Flow A: Generate a customer usage and follow-up workbook."""

import csv
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[3]
INPUT_DIR = ROOT / "artifacts" / "flow_a" / "realdata_downloads"
OUTPUT_DIR = ROOT / "artifacts" / "flow_a"
OUTPUT_FILE = OUTPUT_DIR / "customer_usage_followup_report.xlsx"

EXPECTED_ROWS = {
    "company-a.csv": 5,
    "company-b.csv": 9,
    "company-c.csv": 4,
}

REQUIRED_FIELDS = {
    "company_name",
    "sales",
    "user_id",
    "frequency_label",
}


def read_csv(path):
    """Read one collected CSV and calculate descriptive metrics."""
    with path.open("r", encoding="utf-8-sig", newline="") as file:
        reader = csv.DictReader(file)
        headers = reader.fieldnames or []
        rows = list(reader)

    if not REQUIRED_FIELDS.issubset(set(headers)):
        raise ValueError(f"{path.name}: CSV 缺少必要字段")

    expected = EXPECTED_ROWS.get(path.name)
    if expected is None:
        raise ValueError(f"发现非预期文件：{path.name}")

    if len(rows) != expected:
        raise ValueError(
            f"{path.name}: 预期 {expected} 个账号，实际 {len(rows)} 个"
        )

    date_columns = [
        header for header in headers
        if header not in REQUIRED_FIELDS
    ]

    accounts = []

    for row in rows:
        values = []
        active_dates = []
        missing_days = 0
        invalid_days = 0

        for day in date_columns:
            raw = (row.get(day) or "").strip()

            if not raw:
                missing_days += 1
                continue

            try:
                number = float(raw)
                if (
                    not number.is_integer()
                    or number < 0
                ):
                    raise ValueError
                usage = int(number)
            except ValueError:
                invalid_days += 1
                continue

            values.append(usage)

            if usage > 0:
                active_dates.append(day)

        if not values:
            total_usage = None
            account_status = "数据缺失，暂不判断"
        else:
            total_usage = sum(values)
            account_status = (
                "存在无效数据，人工核查"
                if invalid_days
                else "数据可用"
            )

        label = (row.get("frequency_label") or "").strip()

        if account_status != "数据可用":
            recommendation = account_status
        elif label == "低频":
            recommendation = "优先检查使用情况，结合客户背景判断是否跟进"
        elif label == "高频":
            recommendation = "常规观察，关注后续使用变化"
        else:
            recommendation = "原平台标签不明确，人工核查"

        last_active = max(active_dates) if active_dates else ""

        accounts.append({
            "公司": row["company_name"],
            "负责商务": row["sales"],
            "用户ID": row["user_id"],
            "原平台高低频标签": label,
            "累计使用量": total_usage,
            "有效日期数": len(values),
            "缺失日期数": missing_days,
            "无效数据数": invalid_days,
            "活跃天数": len(active_dates),
            "最近活跃日期": last_active,
            "数据状态": account_status,
            "跟进建议": recommendation,
        })

    return accounts, len(date_columns)


def style_sheet(ws):
    """Apply basic formatting and readable column widths."""
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="34495E",
        )
        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True,
        )

    ws.row_dimensions[1].height = 32

    for column_cells in ws.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_length = max(
            len(str(cell.value or ""))
            for cell in column_cells
        )
        ws.column_dimensions[column_letter].width = min(
            max(max_length + 2, 12), 48
        )

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )


def main():
    if not INPUT_DIR.exists():
        raise FileNotFoundError(f"下载目录不存在：{INPUT_DIR}")

    all_accounts = []
    company_date_counts = {}

    for filename in EXPECTED_ROWS:
        path = INPUT_DIR / filename

        if not path.exists():
            raise FileNotFoundError(f"缺少采集文件：{filename}")

        accounts, date_count = read_csv(path)
        all_accounts.extend(accounts)

        company = accounts[0]["公司"] if accounts else filename
        company_date_counts[company] = date_count

        print(
            f"{company}: {len(accounts)} 个账号，"
            f"{date_count} 个日期列"
        )

    if len(all_accounts) != 18:
        raise ValueError(
            f"预期 18 个账号，实际 {len(all_accounts)} 个"
        )

    # Company-level summary.
    grouped = defaultdict(list)
    for account in all_accounts:
        grouped[account["公司"]].append(account)

    company_rows = []

    for company, accounts in grouped.items():
        company_rows.append({
            "公司": company,
            "负责商务": accounts[0]["负责商务"],
            "账号总数": len(accounts),
            "高频账号数": sum(
                item["原平台高低频标签"] == "高频"
                for item in accounts
            ),
            "低频账号数": sum(
                item["原平台高低频标签"] == "低频"
                for item in accounts
            ),
            "数据缺失账号数": sum(
                item["数据状态"] == "数据缺失，暂不判断"
                for item in accounts
            ),
            "日期列数": company_date_counts[company],
        })

    # Sort low-frequency accounts first, without recalculating labels.
    label_order = {"低频": 0, "高频": 1}
    all_accounts.sort(
        key=lambda item: (
            item["负责商务"],
            label_order.get(item["原平台高低频标签"], 2),
            item["用户ID"],
        )
    )

    wb = Workbook()

    summary_ws = wb.active
    summary_ws.title = "公司汇总"
    summary_headers = list(company_rows[0].keys())
    summary_ws.append(summary_headers)

    for row in company_rows:
        summary_ws.append([
            row[header] for header in summary_headers
        ])

    summary_ws.append([])
    summary_ws.append(["总计", "", len(all_accounts)])
    summary_ws.append([
        "说明",
        "高低频标签沿用原平台结果；不重新计算阈值。",
    ])
    summary_ws.append([
        "说明",
        "缺失数据不按零计算；不据此直接判定客户不健康。",
    ])

    detail_ws = wb.create_sheet("账号跟进明细")
    detail_headers = list(all_accounts[0].keys())
    detail_ws.append(detail_headers)

    for account in all_accounts:
        detail_ws.append([
            account[header] for header in detail_headers
        ])

    style_sheet(summary_ws)
    style_sheet(detail_ws)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_FILE)

    # Re-open the generated workbook to verify its structure.
    check_wb = load_workbook(OUTPUT_FILE, read_only=True, data_only=True)
    try:
        assert check_wb.sheetnames == ["公司汇总", "账号跟进明细"]
        assert check_wb["账号跟进明细"].max_row == 19
        assert check_wb["公司汇总"].max_row >= 4
    finally:
        check_wb.close()

    print("\n===== 跟进报告生成成功 =====")
    print(f"公司数：{len(company_rows)}")
    print(f"账号数：{len(all_accounts)}")
    print(f"报告路径：{OUTPUT_FILE}")
    print("工作表：公司汇总、账号跟进明细")
    print("邮箱未纳入报告。")


if __name__ == "__main__":
    main()
