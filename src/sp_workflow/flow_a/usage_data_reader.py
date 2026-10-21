
"""Flow A: Read and validate SocialPeta usage exports.

Real customer emails are deliberately excluded from all outputs.
Missing values remain missing; they are never silently converted to zero.
"""

import csv
import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[3]
INPUT_DIR = PROJECT_ROOT / "data" / "flow_a_demo"
OUTPUT_DIR = PROJECT_ROOT / "artifacts" / "flow_a"
EXPECTED_FILES = 3


def parse_date(value):
    """Return an ISO date string, or None if this is not a date header."""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()

    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
            try:
                return datetime.strptime(text, fmt).date().isoformat()
            except ValueError:
                pass

    return None


def read_metadata(ws):
    """Read company metadata from the top section of the worksheet."""
    metadata = {}

    for row in ws.iter_rows(min_row=1, max_row=6, values_only=True):
        if len(row) >= 2 and row[0] is not None:
            key = str(row[0]).strip()
            metadata[key] = row[1]

    return {
        "company_id": metadata.get("公司id"),
        "company_name": metadata.get("公司名", "未命名公司"),
        "report_period": metadata.get("时间"),
        "module": metadata.get("模块"),
    }


def read_workbook(path):
    """Normalize one workbook into account summaries and daily records."""
    wb = load_workbook(path, read_only=True, data_only=True)

    try:
        if "访问数据" not in wb.sheetnames:
            raise ValueError("缺少工作表：访问数据")

        ws = wb["访问数据"]
        metadata = read_metadata(ws)

        header_row = None
        headers = []

        for row_idx, row in enumerate(
            ws.iter_rows(values_only=True), start=1
        ):
            values = list(row)

            if values and str(values[0]).strip() == "用户id":
                header_row = row_idx
                headers = values
                break

        if header_row is None:
            raise ValueError("找不到用户数据表头")

        required = {"用户id", "用户邮箱", "高低频率"}
        header_positions = {}

        for index, value in enumerate(headers):
            if value is not None:
                header_positions[str(value).strip()] = index

        missing_headers = required - set(header_positions)
        if missing_headers:
            raise ValueError(
                f"缺少必要字段：{sorted(missing_headers)}"
            )

        user_col = header_positions["用户id"]
        label_col = header_positions["高低频率"]

        date_columns = []
        for index, value in enumerate(headers):
            parsed = parse_date(value)
            if parsed is not None:
                date_columns.append((index, parsed))

        if not date_columns:
            raise ValueError("没有识别到任何日期列")

        accounts = []
        daily_records = []
        seen_user_ids = set()
        duplicate_user_ids = []
        all_empty_accounts = 0

        for row in ws.iter_rows(
            min_row=header_row + 1,
            values_only=True,
        ):
            values = list(row)

            if len(values) <= user_col or values[user_col] is None:
                continue

            user_id = str(values[user_col]).strip()
            if not user_id:
                continue

            if user_id in seen_user_ids:
                duplicate_user_ids.append(user_id)
            seen_user_ids.add(user_id)

            label = values[label_col] if label_col < len(values) else None
            label = str(label).strip() if label is not None else ""

            numeric_values = []
            missing_count = 0
            invalid_count = 0
            active_dates = []
            valid_dates = []

            for col_idx, day in date_columns:
                raw = values[col_idx] if col_idx < len(values) else None

                if raw is None or (
                    isinstance(raw, str) and not raw.strip()
                ):
                    missing_count += 1
                    status = "missing"
                    usage = None

                elif isinstance(raw, bool):
                    invalid_count += 1
                    status = "invalid"
                    usage = None

                elif isinstance(raw, (int, float)):
                    if raw < 0 or (
                        isinstance(raw, float) and not raw.is_integer()
                    ):
                        invalid_count += 1
                        status = "invalid"
                        usage = None
                    else:
                        usage = int(raw)
                        numeric_values.append(usage)
                        valid_dates.append(day)

                        if usage > 0:
                            active_dates.append(day)
                            status = "positive"
                        else:
                            status = "zero"

                else:
                    invalid_count += 1
                    status = "invalid"
                    usage = None

                daily_records.append({
                    "company_name": metadata["company_name"],
                    "user_id": user_id,
                    "frequency_label": label,
                    "date": day,
                    "usage_count": usage,
                    "value_status": status,
                })

            total_usage = sum(numeric_values)
            all_missing = len(numeric_values) == 0

            if all_missing:
                all_empty_accounts += 1

            accounts.append({
                "company_name": metadata["company_name"],
                "user_id": user_id,
                "frequency_label": label,
                "date_columns": len(date_columns),
                "valid_days": len(numeric_values),
                "missing_days": missing_count,
                "invalid_days": invalid_count,
                "total_usage": total_usage if not all_missing else None,
                "active_days": len(active_dates),
                "max_daily_usage": (
                    max(numeric_values) if numeric_values else None
                ),
                "last_active_date": (
                    max(active_dates) if active_dates else None
                ),
                "all_usage_missing": all_missing,
            })

        dates = [day for _, day in date_columns]

        report = {
            "source_file": path.name,
            **metadata,
            "account_count": len(accounts),
            "date_column_count": len(date_columns),
            "first_date_column": min(dates) if dates else None,
            "last_date_column": max(dates) if dates else None,
            "date_columns": dates,
            "all_empty_accounts": all_empty_accounts,
            "duplicate_user_ids": sorted(set(duplicate_user_ids)),
            "accounts": accounts,
        }

        return report, daily_records

    finally:
        wb.close()


def main():
    if not INPUT_DIR.exists():
        raise FileNotFoundError(f"输入目录不存在：{INPUT_DIR}")

    files = sorted(INPUT_DIR.glob("*.xlsx"))

    if len(files) != EXPECTED_FILES:
        raise ValueError(
            f"预期 {EXPECTED_FILES} 份 Excel，实际找到 {len(files)} 份。"
            "请检查输入目录，不要混入其他文件。"
        )

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    reports = []
    all_accounts = []
    all_daily_records = []

    for path in files:
        report, daily_records = read_workbook(path)
        reports.append(report)
        all_accounts.extend(report["accounts"])
        all_daily_records.extend(daily_records)

        print(f"\n文件：{path.name}")
        print(f"公司：{report['company_name']}")
        print(f"账号数量：{report['account_count']}")
        print(f"日期列数量：{report['date_column_count']}")
        print(
            f"日期列范围：{report['first_date_column']} "
            f"至 {report['last_date_column']}"
        )
        print(f"全日期均缺失的账号：{report['all_empty_accounts']}")
        print(f"重复账号 ID：{report['duplicate_user_ids']}")

    # Exclude emails from every output file.
    with (OUTPUT_DIR / "usage_data_quality_report.json").open(
        "w", encoding="utf-8"
    ) as file:
        json.dump(
            {
                "source_file_count": len(files),
                "account_count": len(all_accounts),
                "daily_record_count": len(all_daily_records),
                "companies": reports,
            },
            file,
            ensure_ascii=False,
            indent=2,
        )

    with (OUTPUT_DIR / "usage_account_summary.csv").open(
        "w", encoding="utf-8-sig", newline=""
    ) as file:
        if all_accounts:
            writer = csv.DictWriter(
                file,
                fieldnames=list(all_accounts[0].keys()),
            )
            writer.writeheader()
            writer.writerows(all_accounts)

    with (OUTPUT_DIR / "usage_daily_records.csv").open(
        "w", encoding="utf-8-sig", newline=""
    ) as file:
        if all_daily_records:
            writer = csv.DictWriter(
                file,
                fieldnames=list(all_daily_records[0].keys()),
            )
            writer.writeheader()
            writer.writerows(all_daily_records)

    print("\n===== 测试汇总 =====")
    print(f"文件数：{len(reports)}")
    print(f"账号数：{len(all_accounts)}")
    print(f"每日记录数：{len(all_daily_records)}")
    print(f"输出目录：{OUTPUT_DIR}")
    print("已生成：")
    print("- usage_data_quality_report.json")
    print("- usage_account_summary.csv")
    print("- usage_daily_records.csv")
    print("注意：所有输出均不包含用户邮箱。")


if __name__ == "__main__":
    main()
